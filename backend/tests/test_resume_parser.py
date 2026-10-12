"""resume_parser.py 纯函数单元测试。

覆盖：纯文本（UTF-8 / GB18030 / 未知扩展名）、DOCX（python-docx 生成样例）、
PDF 降级（损坏字节不抛异常）。
"""
import io

import pytest

from app.services.resume_parser import parse_resume


def test_parse_txt_utf8():
    text = "张三\n5 年 Python 后端经验，熟悉 FastAPI 与 PostgreSQL。"
    assert parse_resume(text.encode("utf-8"), "resume.txt") == text


def test_parse_txt_gb18030():
    text = "李四\nJava 工程师，熟悉微服务架构。"
    assert parse_resume(text.encode("gb18030"), "resume.txt") == text


def test_parse_md_and_unknown_extension_fallback_to_text():
    text = "王五\n全栈工程师。"
    assert parse_resume(text.encode("utf-8"), "resume.md") == text
    assert parse_resume(text.encode("utf-8"), "resume.xyz") == text
    assert parse_resume(text.encode("utf-8"), "") == text


def test_parse_empty_content():
    assert parse_resume(b"", "resume.txt") == ""


def test_parse_docx():
    """用 python-docx 生成样例 DOCX，验证段落抽取。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("简历标题：候选人简历")
    doc.add_paragraph("工作经历：某公司后端工程师")
    buf = io.BytesIO()
    doc.save(buf)

    text = parse_resume(buf.getvalue(), "resume.docx")
    assert "候选人简历" in text
    assert "后端工程师" in text


def test_parse_pdf_corrupt_falls_back_no_exception():
    """损坏 PDF 字节应走降级路径（不抛异常）。"""
    result = parse_resume(b"%PDF-1.4 not-a-real-pdf", "resume.pdf")
    # 降级为文本解码，能解出多少算多少，但绝不抛异常
    assert isinstance(result, str)
